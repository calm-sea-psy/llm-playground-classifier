"""KORIE 필드 정답 라벨 사람 검수용 엑셀 (todo C: 라벨 오류율 측정).

정답 라벨(data/verify/korie_fields_labels.csv)은 사람이 아니라 Claude 가 이미지를 보고 만들었다.
표본을 사람이 확인해 라벨 오류율을 재고, 틀린 라벨은 고친 뒤 모델 선정 수치를 다시 채점한다.

표본: OCR 초안을 Claude 가 고친 행(edited=1) 20건 + 그대로 둔 행 10건, 필드 종류가 고르게 섞이도록 뽑음.
OCR 초안 값은 보여 주지 않는다 (검수자가 라벨과 OCR 중 하나를 고르는 식으로 치우치지 않게).

예) src/OcrService/.venv/Scripts/python src/OcrService/tools/label_review.py
결과: data/verify/korie_label_review.xlsx ➔ 판정 열을 채워 저장하면 eval/ 에서 읽어 오류율 계산
"""

import argparse
import csv
import io
import random
import sys
from collections import defaultdict
from pathlib import Path

import cv2
import numpy as np
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image

SERVICE_DIR = Path(__file__).resolve().parent.parent
REPO_ROOT = SERVICE_DIR.parent.parent
sys.path.insert(0, str(SERVICE_DIR))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from imaging import read_image  # noqa: E402
from korie_fields import read_labels  # noqa: E402

FIELD_NAMES = {
    "store_name": "상호", "date": "날짜", "time": "시각", "receipt_no": "영수증 번호",
    "subtotal": "소계(공급가액)", "tax": "세금", "total": "합계", "phone": "전화", "address": "주소",
}
VERDICTS = ["맞음", "틀림", "애매함"]
HEADERS = ["No", "영수증", "필드", "필드 영역", "주변 포함", "라벨 (정답 후보)", "판정", "올바른 값 (틀림일 때)", "비고"]
WIDTHS = [5, 11, 13, 50, 68, 30, 9, 26, 24]
ROW_HEIGHT_PT = 118
FIELD_CROP = (60, 360)      # 높이, 최대 너비 (px)
CONTEXT_CROP = (150, 480)
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")


def crop(image: np.ndarray, box: tuple[int, int, int, int], height: int, max_width: int) -> io.BytesIO:
    x0, y0, x1, y1 = box
    part = Image.fromarray(cv2.cvtColor(image[y0:y1, x0:x1], cv2.COLOR_BGR2RGB))
    scale = min(height / part.height, max_width / part.width)
    part = part.resize((max(int(part.width * scale), 1), max(int(part.height * scale), 1)))
    buffer = io.BytesIO()
    part.save(buffer, format="PNG")
    buffer.seek(0)
    return buffer


def context_box(bbox, width: int, height: int) -> tuple[int, int, int, int]:
    """필드 영역 위아래로 한 칸씩, 좌우로 넓혀 잘린 글자와 옆 칸을 함께 보이게 함"""
    (x0, y0), (x1, y1) = bbox[0], bbox[2]
    h = y1 - y0
    return (max(x0 - 80, 0), max(y0 - int(h * 1.5), 0), min(x1 + 80, width), min(y1 + int(h * 1.5), height))


def sample(rows: list[dict], edited: int, unedited: int, seed: int) -> list[dict]:
    rng = random.Random(seed)
    picked = []
    for flag, count in (("1", edited), ("0", unedited)):
        by_field = defaultdict(list)
        for r in rows:
            if r["edited"] == flag:
                by_field[r["field"]].append(r)
        for group in by_field.values():
            rng.shuffle(group)
        # 필드별로 돌아가며 한 건씩 ➔ 필드 종류가 고르게
        fields = sorted(by_field)
        while count > 0 and any(by_field.values()):
            for f in fields:
                if count > 0 and by_field[f]:
                    picked.append(by_field[f].pop())
                    count -= 1
    rng.shuffle(picked)
    return picked


def main():
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--labels", type=Path, default=REPO_ROOT / "data" / "verify" / "korie_fields_labels.csv")
    p.add_argument("--korie", type=Path, default=REPO_ROOT / "data" / "samples" / "korie")
    p.add_argument("--out", type=Path, default=REPO_ROOT / "data" / "verify" / "korie_label_review.xlsx")
    p.add_argument("--edited", type=int, default=20)
    p.add_argument("--unedited", type=int, default=10)
    p.add_argument("--seed", type=int, default=7)
    args = p.parse_args()

    rows = [r for r in csv.DictReader(open(args.labels, encoding="utf-8-sig"))
            if r["label"] and r["label"] != "#INVALID" and r["field"] in FIELD_NAMES]
    picked = sample(rows, args.edited, args.unedited, args.seed)
    images = {p.stem: p for p in (args.korie / "images").iterdir()}

    wb = Workbook()
    ws = wb.active
    ws.title = "검수"
    ws.append(HEADERS)
    for i, width in enumerate(WIDTHS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"
    verdict = DataValidation(type="list", formula1='"' + ",".join(VERDICTS) + '"', allow_blank=True)
    ws.add_data_validation(verdict)

    for n, row in enumerate(picked, start=1):
        image = read_image(images[row["image"]])
        h, w = image.shape[:2]
        bbox = read_labels(args.korie / "labels" / f"{row['image']}.txt", w, h)[row["field"]]
        r = n + 1
        ws.append([n, row["image"], FIELD_NAMES[row["field"]], None, None, row["label"], None, None, None])
        ws.row_dimensions[r].height = ROW_HEIGHT_PT
        for col in range(1, 10):
            ws.cell(row=r, column=col).alignment = Alignment(vertical="center", wrap_text=True)
        ws.cell(row=r, column=6).number_format = "@"
        ws.cell(row=r, column=6).font = Font(size=13, bold=True)
        ws.cell(row=r, column=8).number_format = "@"
        verdict.add(f"G{r}")
        (x0, y0), (x1, y1) = bbox[0], bbox[2]
        ws.add_image(XlImage(crop(image, (x0, y0, x1, y1), *FIELD_CROP)), f"D{r}")
        ws.add_image(XlImage(crop(image, context_box(bbox, w, h), *CONTEXT_CROP)), f"E{r}")
        print(f"[{n}/{len(picked)}] {row['image']} {row['field']}")

    # 채점용 원본 정보 (검수자에게는 숨김)
    meta = wb.create_sheet("meta")
    meta.append(["No", "image", "field", "label", "edited", "ocr_normalized"])
    for n, row in enumerate(picked, start=1):
        meta.append([n, row["image"], row["field"], row["label"], int(row["edited"]), row["ocr_normalized"]])
    meta.sheet_state = "hidden"

    guide = wb.create_sheet("설명", 0)
    for line in [
        "KORIE 영수증 필드 정답 라벨 검수",
        "",
        "이 라벨은 사람이 아니라 AI(Claude)가 영수증 이미지를 보고 만든 것입니다. 표본 30건을 사람이 확인해 라벨 오류율을 잽니다.",
        "",
        "하는 일",
        "1. '검수' 시트에서 '필드 영역'과 '주변 포함' 이미지를 보고, 'F 라벨' 값이 영수증에 인쇄된 값과 같은지 판단합니다.",
        "2. G '판정' 열에서 맞음 / 틀림 / 애매함 중 하나를 고릅니다.",
        "3. 틀림이면 H 열에 올바른 값을 적습니다. 애매함이면 I 비고에 이유(흐림, 잘림 등)를 적어 주세요.",
        "4. 저장 후 알려 주시면 오류율을 계산하고 틀린 라벨을 고쳐 다시 채점합니다.",
        "",
        "판단 기준 (라벨 형식)",
        "- 금액(소계·세금·합계): 숫자만 (쉼표·'원' 없음). 값이 같으면 맞음",
        "- 날짜: YYYY-MM-DD (두 자리 연도는 20을 붙임), 시각: 인쇄된 자리수까지 (HH:MM 또는 HH:MM:SS)",
        "- 상호·주소·번호·전화: 인쇄된 글자 그대로 (띄어쓰기 차이는 무시)",
        "- '필드 영역'은 KORIE 원본 라벨의 박스를 잘라낸 것입니다. 박스가 어긋나 값이 잘려 보이면 '주변 포함' 이미지로 판단하고 비고에 적어 주세요.",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 130
    guide["A1"].font = Font(bold=True, size=14)
    wb.active = 1

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"완료 ➔ {args.out}")


if __name__ == "__main__":
    main()
