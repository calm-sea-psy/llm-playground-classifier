"""검증용 엑셀 (todo 1번 ㄹ)-g): 사람이 OCR 결과를 원본과 대조해 정답 열을 채우는 용도.

lines 시트의 H(정답 텍스트) 열은 eval/ 스크립트가 CER 라벨로 다시 읽는다.
열 순서를 바꾸면 eval 쪽도 같이 고쳐야 한다.
"""

import io
import json

import cv2
import numpy as np
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image

from contract import OcrResult

LOW_CONFIDENCE = 0.8
ERROR_TYPES = ["오인식", "누락", "과검출(없는 글자)", "줄 분리·병합", "순서"]

LINE_HEADERS = ["문서명", "페이지", "줄번호", "OCR 텍스트", "신뢰도", "bbox",
                "줄 이미지", "정답 텍스트", "일치", "오류 유형", "비고"]
LINE_WIDTHS = [34, 7, 7, 40, 8, 22, 52, 40, 6, 16, 20]

SUMMARY_HEADERS = ["문서명", "엔진", "모델", "처리시간(ms)", "줄 수", "평균 신뢰도",
                   "일치 줄 수", "일치율", f"신뢰도<{LOW_CONFIDENCE} 줄 수", *ERROR_TYPES]
SUMMARY_WIDTHS = [34, 11, 44, 12, 7, 11, 10, 8, 14, *[12] * len(ERROR_TYPES)]

CROP_HEIGHT_PX = 28
CROP_MAX_WIDTH_PX = 360
ROW_HEIGHT_PT = 24

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
RED = PatternFill("solid", fgColor="F8CBAD")


def _setup_sheet(ws, headers, widths):
    ws.append(headers)
    for i, width in enumerate(widths, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"


def _crop_line(image: np.ndarray, bbox: list[list[int]], height_px: int = CROP_HEIGHT_PX) -> io.BytesIO:
    h, w = image.shape[:2]
    xs = [p[0] for p in bbox]
    ys = [p[1] for p in bbox]
    x0, x1 = max(min(xs) - 2, 0), min(max(xs) + 2, w)
    y0, y1 = max(min(ys) - 2, 0), min(max(ys) + 2, h)
    crop = Image.fromarray(cv2.cvtColor(image[y0:y1, x0:x1], cv2.COLOR_BGR2RGB))
    scale = min(height_px / crop.height, CROP_MAX_WIDTH_PX / crop.width)
    crop = crop.resize((max(int(crop.width * scale), 1), max(int(crop.height * scale), 1)))
    buffer = io.BytesIO()
    crop.save(buffer, format="PNG")
    buffer.seek(0)
    return buffer


class VerifyWorkbook:
    def __init__(self):
        self.wb = Workbook()
        self.lines = self.wb.active
        self.lines.title = "lines"
        self.summary = self.wb.create_sheet("summary")
        _setup_sheet(self.lines, LINE_HEADERS, LINE_WIDTHS)
        _setup_sheet(self.summary, SUMMARY_HEADERS, SUMMARY_WIDTHS)

    def add_document(self, name: str, image: np.ndarray, result: OcrResult):
        ws = self.lines
        for no, line in enumerate(result.lines, start=1):
            r = ws.max_row + 1
            ws.append([
                name, result.page, no, line.text, line.confidence,
                json.dumps(line.bbox, separators=(",", ":")) if line.bbox else None,
                None,
                line.text,  # 정답 텍스트를 OCR 값으로 미리 채워 두고 사람은 틀린 부분만 고침
                f'=IF(TRIM(D{r})=TRIM(H{r}),"O","X")',
            ])
            ws.cell(row=r, column=4).alignment = Alignment(vertical="center")
            ws.cell(row=r, column=8).alignment = Alignment(vertical="center")
            ws.row_dimensions[r].height = ROW_HEIGHT_PT
            if line.bbox:
                ws.add_image(XlImage(_crop_line(image, line.bbox)), f"G{r}")

        s = self.summary
        r = s.max_row + 1
        name_ref = f"$A{r}"
        s.append([
            name, result.engine, result.model_version, result.elapsed_ms,
            f"=COUNTIF(lines!$A:$A,{name_ref})",
            result.avg_confidence,
            f'=COUNTIFS(lines!$A:$A,{name_ref},lines!$I:$I,"O")',
            f"=IFERROR(G{r}/E{r},0)",
            f'=COUNTIFS(lines!$A:$A,{name_ref},lines!$E:$E,"<{LOW_CONFIDENCE}")',
            *[f"=COUNTIFS(lines!$A:$A,{name_ref},lines!$J:$J,{col}$1)"
              for col in ("J", "K", "L", "M", "N")],
        ])
        s.cell(row=r, column=8).number_format = "0.0%"

    def _finish(self):
        ws, s = self.lines, self.summary
        last = max(ws.max_row, 2)

        # 오류 유형 드롭다운
        dv = DataValidation(type="list", formula1=f'"{",".join(ERROR_TYPES)}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"J2:J{last}")

        # 신뢰도 낮은 줄은 노란색(우선 확인), 불일치는 빨간색
        ws.conditional_formatting.add(
            f"D2:E{last}", FormulaRule(formula=[f"AND(ISNUMBER($E2),$E2<{LOW_CONFIDENCE})"], fill=YELLOW))
        ws.conditional_formatting.add(f"H2:I{last}", FormulaRule(formula=['$I2="X"'], fill=RED))
        ws.auto_filter.ref = f"A1:K{last}"

        # summary 전체 합계 행
        first, end = 2, s.max_row
        if end >= first:
            t = end + 1
            s.append([
                "전체", None, None,
                f"=SUM(D{first}:D{end})",
                f"=SUM(E{first}:E{end})",
                "=IFERROR(AVERAGE(lines!E:E),0)",
                f"=SUM(G{first}:G{end})",
                f"=IFERROR(G{t}/E{t},0)",
                f"=SUM(I{first}:I{end})",
                *[f"=SUM({c}{first}:{c}{end})" for c in ("J", "K", "L", "M", "N")],
            ])
            for cell in s[t]:
                cell.font = HEADER_FONT
            s.cell(row=t, column=8).number_format = "0.0%"
            s.cell(row=t, column=6).number_format = "0.0000"

    def save(self, target):
        self._finish()
        self.wb.save(target)


FIELD_HEADERS = ["문서명", "필드명", "OCR 원문", "정규화값", "필드 이미지", "정답", "일치", "비고"]
FIELD_WIDTHS = [18, 14, 40, 30, 52, 30, 6, 20]
FIELD_SUMMARY_HEADERS = ["필드명", "행 수", "일치 수", "일치율"]
# 주소처럼 여러 줄인 필드도 읽을 수 있게 줄 이미지보다 크게
FIELD_CROP_HEIGHT_PX = 44
FIELD_ROW_HEIGHT_PT = 36


class FieldVerifyWorkbook:
    """필드 단위 검증 엑셀 (fields 시트). 정답(F) 열은 eval/ 에서 필드 정답 라벨로 읽는다."""

    def __init__(self, field_names: list[str]):
        self.field_names = field_names
        self.wb = Workbook()
        self.fields = self.wb.active
        self.fields.title = "fields"
        self.summary = self.wb.create_sheet("summary")
        _setup_sheet(self.fields, FIELD_HEADERS, FIELD_WIDTHS)
        _setup_sheet(self.summary, FIELD_SUMMARY_HEADERS, [16, 8, 8, 8])

    def add_field(self, doc: str, field: str, raw: str, normalized: str,
                  image: np.ndarray | None = None, bbox: list[list[int]] | None = None):
        ws = self.fields
        r = ws.max_row + 1
        ws.append([doc, field, raw, normalized, None,
                   normalized,  # 정답은 정규화값으로 미리 채우고 사람은 틀린 것만 고침
                   f'=IF(TRIM(D{r})=TRIM(F{r}),"O","X")'])
        for col in (3, 4, 6):
            cell = ws.cell(row=r, column=col)
            cell.alignment = Alignment(vertical="center")
            cell.number_format = "@"  # 금액·번호가 숫자로 바뀌지 않도록 텍스트로 유지
        ws.row_dimensions[r].height = FIELD_ROW_HEIGHT_PT
        if image is not None and bbox:
            ws.add_image(XlImage(_crop_line(image, bbox, FIELD_CROP_HEIGHT_PX)), f"E{r}")

    def save(self, target):
        ws, s = self.fields, self.summary
        last = max(ws.max_row, 2)
        ws.conditional_formatting.add(f"F2:G{last}", FormulaRule(formula=['$G2="X"'], fill=RED))
        ws.auto_filter.ref = f"A1:H{last}"
        for field in self.field_names:
            r = s.max_row + 1
            s.append([field,
                      f"=COUNTIF(fields!$B:$B,$A{r})",
                      f'=COUNTIFS(fields!$B:$B,$A{r},fields!$G:$G,"O")',
                      f"=IFERROR(C{r}/B{r},0)"])
            s.cell(row=r, column=4).number_format = "0.0%"
        self.wb.save(target)
