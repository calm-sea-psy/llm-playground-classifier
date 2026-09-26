"""4차 0단계: 합성 이력서 생성기 (가짜 인물, 실제 개인정보 없음).

정답 JSON 을 먼저 만들고 같은 내용을 양식 4종 × 형식 3종으로 채운다 ➔ 라벨링이 필요 없다.
- 양식: t1 자유 양식 · t2 표 양식(주민번호 · 가족 사항 · 혼인 여부 포함 ➔ 수집 금지 값 누출 시험) · t3 채용 사이트형 · t4 경력 많은 2쪽
- 형식: doc.pdf (텍스트 층) · doc.docx · scan_N.png (이미지로 그린 뒤 기울기 · 흐림 · 잡음)
- 정답은 문서 종류 팩(packs/resume/type.json) 필드 이름을 그대로 씀

출력: data/digitizer/resume/{split}/{id}/ (git 제외)

예) eval/digitizer/.venv/Scripts/python eval/digitizer/gen_resume.py --split dev --count 24 --seed 1
"""

import argparse
import json
import random
import sys
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

REPO = Path(__file__).resolve().parents[2]
OUT = REPO / "data" / "digitizer" / "resume"
FONT = "C:/Windows/Fonts/malgun.ttf"
FONT_BOLD = "C:/Windows/Fonts/malgunbd.ttf"

# ── 가짜 인물 ────────────────────────────────────────────────────────────
SURNAMES = list("김이박최정강조윤장임한오서신권황안송류홍")
GIVEN = list("민서지하윤도현수아준우예은주원채진태경나영")
ROMAN = {"김": "kim", "이": "lee", "박": "park", "최": "choi", "정": "jung", "강": "kang", "조": "cho", "윤": "yoon",
         "장": "jang", "임": "lim", "한": "han", "오": "oh", "서": "seo", "신": "shin", "권": "kwon", "황": "hwang",
         "안": "ahn", "송": "song", "류": "ryu", "홍": "hong"}
CITIES = [("서울특별시", ["마포구", "강남구", "송파구", "관악구", "노원구"]), ("경기도", ["성남시 분당구", "수원시 영통구", "고양시 일산동구"]),
          ("부산광역시", ["해운대구", "수영구"]), ("충청북도", ["청주시 흥덕구", "청주시 상당구"]), ("대전광역시", ["유성구", "서구"])]
ROADS = ["누리로", "가온길", "한결로", "미리내길", "새봄로", "별빛길", "푸른솔로", "다온대로"]
PREFIX = ["누리", "가온", "다온", "별빛", "새봄", "푸른솔", "한결", "미리내", "아라", "온새미", "이음", "나래"]
COMPANY_SUFFIX = ["테크", "물산", "정보", "시스템즈", "유통", "디자인", "바이오", "로지스", "솔루션", "커머스"]
DEPTS = ["개발팀", "기획팀", "영업1팀", "품질관리팀", "인사팀", "데이터팀", "마케팅팀", "구매팀", "재무팀", "고객지원팀"]
POSITIONS = ["사원", "주임", "대리", "과장", "선임", "책임", "팀장"]
DUTIES = ["사내 주문 관리 시스템 유지보수 및 신규 기능 개발", "월별 매출 보고서 작성과 거래처 관리", "제품 출하 전 품질 검사 기준 수립",
          "채용 공고 운영 및 면접 일정 관리", "고객 데이터 분석 대시보드 구축", "신제품 온라인 광고 캠페인 운영",
          "원자재 구매 단가 협상과 발주 관리", "결산 자료 작성 및 세무 신고 지원", "고객 문의 응대 매뉴얼 개선", "물류 창고 재고 관리 프로세스 개선"]
MAJORS = ["컴퓨터공학", "경영학", "산업공학", "전자공학", "통계학", "화학공학", "국제통상학", "시각디자인", "회계학", "기계공학"]
CERTS = [("정보처리기사", "한국산업인력공단"), ("SQLD", "한국데이터산업진흥원"), ("컴퓨터활용능력 1급", "대한상공회의소"),
         ("전산회계 1급", "한국세무사회"), ("물류관리사", "한국산업인력공단"), ("품질경영기사", "한국산업인력공단"), ("ADsP", "한국데이터산업진흥원")]
LANGS = [("TOEIC", lambda r: str(r.randrange(700, 990, 5))), ("OPIc", lambda r: r.choice(["IM2", "IM3", "IH", "AL"])),
         ("JLPT", lambda r: r.choice(["N1", "N2", "N3"])), ("TOEIC Speaking", lambda r: r.choice(["Level 6", "Level 7", "IH"]))]
SKILLS = ["Python", "Java", "SQL", "Excel", "PowerPoint", "Figma", "C#", "React", "SAP", "Tableau", "Photoshop", "AutoCAD", "R", "Git"]


def ym(y, m):
    return f"{y:04d}-{m:02d}"


def make_person(rng: random.Random, many_careers: bool) -> dict:
    surname = rng.choice(SURNAMES)
    given = "".join(rng.sample(GIVEN, 2))
    birth_y = rng.randint(1978, 1999)
    birth = f"{birth_y:04d}-{rng.randint(1, 12):02d}-{rng.randint(1, 28):02d}"
    city, gus = rng.choice(CITIES)
    address = f"{city} {rng.choice(gus)} {rng.choice(ROADS)} {rng.randint(1, 300)}, {rng.randint(101, 115)}동 {rng.randint(101, 2004)}호"
    email = f"{ROMAN[surname]}.{rng.randint(10, 999)}{rng.choice(['a', 'k', 'j', 's'])}@example.com"
    phone = f"010-{rng.randint(2000, 9999)}-{rng.randint(1000, 9999)}"

    # 학력: 고등학교(가끔) ➔ 대학 ➔ 대학원(가끔)
    education = []
    y = birth_y + 19
    if rng.random() < 0.3:
        education.append({"school": f"{rng.choice(PREFIX)}고등학교", "major": None, "degree": "졸업",  # v0.1 은 "고등학교 졸업" ➔ 원문이 "OO고등학교 고등학교 졸업" 이 되어 전공 · 학위 경계가 모호했음
                          "start": ym(y - 3, 3), "end": ym(y, 2)})
    major = rng.choice(MAJORS)
    education.append({"school": f"{rng.choice(PREFIX)}대학교", "major": major, "degree": "학사",
                      "start": ym(y, 3), "end": ym(y + 4, 2)})
    y += 4
    if rng.random() < 0.25:
        education.append({"school": f"{rng.choice(PREFIX)}대학교 대학원", "major": major, "degree": "석사",
                          "start": ym(y, 3), "end": ym(y + 2, 2)})
        y += 2

    # 경력: 겹치지 않게 이어 붙이고 마지막은 가끔 재직 중
    n = rng.randint(5, 8) if many_careers else rng.randint(1, 3)
    career, cy, cm = [], y, rng.choice([3, 4, 7, 9])
    used = set()
    for i in range(n):
        while True:
            company = rng.choice(PREFIX) + rng.choice(COMPANY_SUFFIX)
            if company not in used:
                used.add(company)
                break
        months = rng.randint(8, 40) if many_careers else rng.randint(14, 60)
        ey, em = cy + (cm - 1 + months) // 12, (cm - 1 + months) % 12 + 1
        if ey > 2026 or (ey == 2026 and em > 8):
            ey, em = 2026, 8
        present = i == n - 1 and rng.random() < 0.5
        career.append({"company": company, "department": rng.choice(DEPTS), "position": POSITIONS[min(i + rng.randint(0, 1), len(POSITIONS) - 1)],
                       "start": ym(cy, cm), "end": "present" if present else ym(ey, em), "duties": rng.choice(DUTIES)})
        cy, cm = (ey, em + 1) if em < 12 else (ey + 1, 1)
        if cy >= 2026:
            break

    certificates = [{"name": c, "issuer": iss, "date": ym(rng.randint(birth_y + 20, 2025), rng.randint(1, 12))}
                    for c, iss in rng.sample(CERTS, rng.randint(0, 3))]
    languages = [{"test": t, "score": f(rng), "date": ym(rng.randint(2018, 2025), rng.randint(1, 12))}
                 for t, f in rng.sample(LANGS, rng.randint(0, 2))]
    skills = rng.sample(SKILLS, rng.randint(2, 5))
    return {"name": surname + given, "phone": phone, "email": email, "address": address, "birth_date": birth,
            "education": education, "career": career, "certificates": certificates, "languages": languages, "skills": skills}


def forbidden_values(rng: random.Random, person: dict) -> dict:
    """표 양식에만 적는 수집 금지 값 (가짜). 추출 결과 어디에도 나오면 안 됨"""
    b = person["birth_date"]
    rid = f"{b[2:4]}{b[5:7]}{b[8:10]}-{rng.choice([1, 2]) if b[:2] == '19' else rng.choice([3, 4])}{rng.randint(100000, 999999)}"
    father = rng.choice(SURNAMES[:1]) + "".join(rng.sample(GIVEN, 2))
    return {"resident_id": rid, "family": [("부", person["name"][0] + "".join(rng.sample(GIVEN, 2)), "자영업"),
                                           ("모", rng.choice(SURNAMES) + "".join(rng.sample(GIVEN, 2)), "교사")],
            "marital_status": rng.choice(["미혼", "기혼"]), "body": f"{rng.randint(155, 188)}cm / {rng.randint(45, 90)}kg",
            "hometown": rng.choice([c for c, _ in CITIES]), "_unused": father}


# ── 날짜 표기 (양식마다 다르게) ─────────────────────────────────────────
def fmt_month(v, style):
    if v is None:
        return ""
    if v == "present":
        return {"dot": "현재", "kor": "재직 중", "dash": "재직중"}[style]
    y, m = v.split("-")
    return {"dot": f"{y}.{m}", "kor": f"{y}년 {int(m)}월", "dash": f"{y}-{m}"}[style]


def fmt_range(a, b, style):
    return f"{fmt_month(a, style)} ~ {fmt_month(b, style)}" if a else fmt_month(b, style)


def fmt_birth(v, style):
    y, m, d = v.split("-")
    return {"dot": f"{y}.{m}.{d}", "kor": f"{y}년 {int(m)}월 {int(d)}일", "dash": v}[style]


# ── 양식: 블록 목록으로 표현 (PDF · 이미지 · DOCX 가 같은 블록을 그림) ────
@dataclass
class Doc:
    blocks: list = field(default_factory=list)

    def add(self, kind, *args):
        self.blocks.append((kind, *args))


ACHIEVEMENTS = ["처리 시간 30% 단축", "신규 거래처 12곳 확보", "불량률 1.2%p 개선", "사내 교육 과정 3개 개설", "월간 보고 자동화",
                "재고 회전율 개선", "고객 만족도 조사 1위 부서 선정", "연간 비용 8천만 원 절감"]


def template_t1(p, st="dot", achievements=None):
    d = Doc()
    d.add("title", "이력서")
    d.add("big", p["name"])
    d.add("lines", [f"연락처: {p['phone']}", f"이메일: {p['email']}", f"주소: {p['address']}", f"생년월일: {fmt_birth(p['birth_date'], st)}"])
    d.add("heading", "학력")
    d.add("bullets", [f"{e['school']} {e['major'] or ''} {e['degree']} ({fmt_range(e['start'], e['end'], st)})".replace("  ", " ") for e in p["education"]])
    d.add("heading", "경력")
    for c in p["career"]:
        d.add("bullets", [f"{c['company']} · {c['department']} · {c['position']} ({fmt_range(c['start'], c['end'], st)})"])
        d.add("para", f"담당 업무: {c['duties']}")
        if achievements:  # t4: 정답에 없는 줄 (2쪽으로 늘리고, 추출하면 안 되는 잡음)
            d.add("bullets", [f"주요 성과: {a}" for a in achievements.sample(ACHIEVEMENTS, 3)])
    if p["certificates"]:
        d.add("heading", "자격증")
        d.add("bullets", [f"{c['name']} ({c['issuer']}, {fmt_month(c['date'], st)})" for c in p["certificates"]])
    if p["languages"]:
        d.add("heading", "어학")
        d.add("bullets", [f"{l['test']} {l['score']} ({fmt_month(l['date'], st)})" for l in p["languages"]])
    d.add("heading", "보유 기술")
    d.add("para", ", ".join(p["skills"]))
    return d


def template_t2(p, fb, st="kor"):
    d = Doc()
    d.add("title", "이 력 서")
    d.add("table", None, [["성명", p["name"], "주민등록번호", fb["resident_id"]],
                          ["생년월일", fmt_birth(p["birth_date"], st), "혼인 여부", fb["marital_status"]],
                          ["휴대전화", p["phone"], "이메일", p["email"]],
                          ["주소", p["address"], "신장/체중", fb["body"]],
                          ["본적", fb["hometown"], "", ""]], [0.15, 0.35, 0.18, 0.32])
    d.add("heading", "학력 사항")
    d.add("table", ["기간", "학교명", "전공", "학위"],
          [[fmt_range(e["start"], e["end"], st), e["school"], e["major"] or "", e["degree"]] for e in p["education"]], [0.3, 0.3, 0.2, 0.2])
    d.add("heading", "경력 사항")
    d.add("table", ["기간", "회사명", "부서", "직위", "담당 업무"],
          [[fmt_range(c["start"], c["end"], st), c["company"], c["department"], c["position"], c["duties"]] for c in p["career"]],
          [0.22, 0.16, 0.13, 0.1, 0.39])
    if p["certificates"] or p["languages"]:
        d.add("heading", "자격 및 어학")
        rows = [[fmt_month(c["date"], st), c["name"], c["issuer"]] for c in p["certificates"]]
        rows += [[fmt_month(l["date"], st), f"{l['test']} {l['score']}", ""] for l in p["languages"]]
        d.add("table", ["취득일", "자격증 · 어학", "발급 기관"], rows, [0.25, 0.45, 0.3])
    d.add("heading", "가족 사항")
    d.add("table", ["관계", "성명", "직업"], [list(f) for f in fb["family"]], [0.2, 0.4, 0.4])
    d.add("heading", "보유 기술")
    d.add("para", " / ".join(p["skills"]))
    return d


def template_t3(p, st="dash"):
    d = Doc()
    d.add("big", p["name"])
    d.add("lines", [f"{p['phone']}  |  {p['email']}", p["address"]])
    d.add("heading", "경력")
    for c in reversed(p["career"]):  # 채용 사이트는 최근 경력이 위
        d.add("lines", [f"{fmt_range(c['start'], c['end'], st)}   {c['company']}", f"{c['department']} / {c['position']}"])
        d.add("bullets", [c["duties"]])
    d.add("heading", "학력")
    for e in reversed(p["education"]):
        # 채용 사이트형은 졸업 연월만 적음 ➔ 정답의 입학은 null
        d.add("lines", [f"{e['school']}  |  {e['major'] or '-'}  |  {e['degree']}  |  졸업 {fmt_month(e['end'], st)}"])
    if p["certificates"]:
        d.add("heading", "자격증")
        d.add("lines", [f"{c['name']}  ·  {c['issuer']}  ·  {fmt_month(c['date'], st)}" for c in p["certificates"]])
    if p["languages"]:
        d.add("heading", "어학")
        d.add("lines", [f"{l['test']}  {l['score']}  ({fmt_month(l['date'], st)})" for l in p["languages"]])
    d.add("heading", "스킬")
    d.add("para", "  ".join(f"#{s}" for s in p["skills"]))
    return d


def truth_for(person, template):
    t = json.loads(json.dumps(person))
    if template == "t3":
        t["birth_date"] = None
        for e in t["education"]:
            e["start"] = None
    return t


# ── 그리기: 공통 배치 엔진 (측정 함수만 다름) ─────────────────────────────
class Layout:
    """A4, 좌표는 pt. draw_text(x, y, text, size, bold) / draw_line(x1, y1, x2, y2) 를 백엔드가 구현"""

    W, H = A4
    M = 50

    def __init__(self, measure, draw_text, draw_line, new_page):
        self.measure, self.text, self.line, self.new_page_cb = measure, draw_text, draw_line, new_page
        self.y = self.M

    def ensure(self, h):
        if self.y + h > self.H - self.M:
            self.new_page_cb()
            self.y = self.M

    def wrap(self, s, size, width, bold=False):
        out, cur = [], ""
        for ch in s:
            if self.measure(cur + ch, size, bold) > width and cur:
                out.append(cur)
                cur = ch.lstrip() if ch == " " else ch
            else:
                cur += ch
        return out + ([cur] if cur else [])

    def render(self, doc: Doc):
        width = self.W - 2 * self.M
        for kind, *a in doc.blocks:
            if kind == "title":
                self.ensure(40); self.text(self.W / 2 - self.measure(a[0], 22, True) / 2, self.y + 22, a[0], 22, True); self.y += 40
            elif kind == "big":
                self.ensure(32); self.text(self.M, self.y + 18, a[0], 18, True); self.y += 30
            elif kind == "heading":
                self.ensure(34); self.y += 8; self.text(self.M, self.y + 13, a[0], 13, True)
                self.line(self.M, self.y + 18, self.W - self.M, self.y + 18); self.y += 26
            elif kind in ("lines", "bullets"):
                for s in a[0]:
                    prefix = "• " if kind == "bullets" else ""
                    for i, part in enumerate(self.wrap(prefix + s, 10.5, width - 10)):
                        self.ensure(16); self.text(self.M + (10 if i else 0), self.y + 11, part, 10.5, False); self.y += 16
            elif kind == "para":
                for part in self.wrap(a[0], 10.5, width - 20):
                    self.ensure(16); self.text(self.M + 12, self.y + 11, part, 10.5, False); self.y += 16
            elif kind == "table":
                header, rows, ratios = a
                cols = [r * width for r in ratios]
                for ri, row in enumerate(([header] if header else []) + rows):
                    cells = [self.wrap(str(c), 9.5, w - 8, ri == 0 and header is not None) for c, w in zip(row, cols)]
                    h = max(len(c) for c in cells) * 14 + 8
                    self.ensure(h)
                    x = self.M
                    self.line(self.M, self.y, self.W - self.M, self.y)
                    for cw, lines in zip(cols, cells):
                        self.line(x, self.y, x, self.y + h)
                        for li, t in enumerate(lines):
                            self.text(x + 4, self.y + 14 + li * 14, t, 9.5, ri == 0 and header is not None)
                        x += cw
                    self.line(self.W - self.M, self.y, self.W - self.M, self.y + h)
                    self.y += h
                self.line(self.M, self.y, self.W - self.M, self.y)
                self.y += 6


def render_pdf(doc: Doc, path: Path):
    c = canvas.Canvas(str(path), pagesize=A4)
    H = A4[1]
    measure = lambda s, size, bold: pdfmetrics.stringWidth(s, "MalgunBd" if bold else "Malgun", size)

    def text(x, y, s, size, bold):
        c.setFont("MalgunBd" if bold else "Malgun", size)
        c.drawString(x, H - y, s)

    def line(x1, y1, x2, y2):
        c.setLineWidth(0.6)
        c.line(x1, H - y1, x2, H - y2)

    Layout(measure, text, line, c.showPage).render(doc)
    c.save()


def render_images(doc: Doc, scale=200 / 72):
    pages = []
    size = (int(A4[0] * scale), int(A4[1] * scale))
    fonts = {}

    def font(sz, bold):
        key = (sz, bold)
        if key not in fonts:
            fonts[key] = ImageFont.truetype(FONT_BOLD if bold else FONT, int(sz * scale))
        return fonts[key]

    def new_page():
        img = Image.new("L", size, 255)
        pages.append((img, ImageDraw.Draw(img)))

    new_page()
    measure = lambda s, sz, bold: font(sz, bold).getlength(s) / scale
    text = lambda x, y, s, sz, bold: pages[-1][1].text((x * scale, (y - sz) * scale), s, fill=0, font=font(sz, bold))
    line = lambda x1, y1, x2, y2: pages[-1][1].line((x1 * scale, y1 * scale, x2 * scale, y2 * scale), fill=0, width=2)
    Layout(measure, text, line, new_page).render(doc)
    return [p[0] for p in pages]


def scan_effect(img: Image.Image, rng: random.Random) -> Image.Image:
    """출력 ➔ 스캔 흉내: 기울기 · 흐림 · 잡음 · 대비 저하"""
    img = img.rotate(rng.uniform(-1.5, 1.5), resample=Image.BICUBIC, expand=False, fillcolor=255)
    img = img.filter(ImageFilter.GaussianBlur(rng.uniform(0.5, 1.1)))
    a = np.asarray(img).astype(np.float32)
    a = a * rng.uniform(0.85, 0.95) + rng.uniform(10, 25)  # 종이가 약간 회색
    a += np.random.default_rng(rng.randint(0, 1 << 30)).normal(0, rng.uniform(4, 9), a.shape)
    return Image.fromarray(np.clip(a, 0, 255).astype(np.uint8))


def render_docx(doc: Doc, path: Path):
    d = Document()
    style = d.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    style.font.size = Pt(10.5)
    for kind, *a in doc.blocks:
        if kind == "title":
            d.add_heading(a[0], level=0)
        elif kind == "big":
            r = d.add_paragraph().add_run(a[0]); r.bold = True; r.font.size = Pt(18)
        elif kind == "heading":
            d.add_heading(a[0], level=2)
        elif kind == "lines":
            for s in a[0]:
                d.add_paragraph(s)
        elif kind == "bullets":
            for s in a[0]:
                d.add_paragraph(s, style="List Bullet")
        elif kind == "para":
            d.add_paragraph(a[0])
        elif kind == "table":
            header, rows, _ = a
            all_rows = ([header] if header else []) + rows
            t = d.add_table(rows=len(all_rows), cols=len(all_rows[0]))
            t.style = "Table Grid"
            for i, row in enumerate(all_rows):
                for j, v in enumerate(row):
                    t.cell(i, j).text = str(v)
    d.save(str(path))


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--split", required=True, choices=["dev", "holdout"])
    ap.add_argument("--count", type=int, required=True)
    ap.add_argument("--seed", type=int, required=True)
    args = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")

    pdfmetrics.registerFont(TTFont("Malgun", FONT))
    pdfmetrics.registerFont(TTFont("MalgunBd", FONT_BOLD))
    rng = random.Random(args.seed)
    templates = ["t1", "t2", "t3", "t4"]
    styles = {"t1": "dot", "t2": "kor", "t3": "dash", "t4": "dot"}
    out_dir = OUT / args.split
    out_dir.mkdir(parents=True, exist_ok=True)

    manifest = []
    for i in range(args.count):
        template = templates[i % len(templates)]
        person = make_person(rng, many_careers=template == "t4")
        fb = forbidden_values(rng, person)
        if template == "t2":
            doc = template_t2(person, fb, styles[template])
        elif template == "t3":
            doc = template_t3(person, styles[template])
        else:
            doc = template_t1(person, styles[template], achievements=rng if template == "t4" else None)
        doc_id = f"{args.split}-{i + 1:03d}-{template}"
        d = out_dir / doc_id
        d.mkdir(exist_ok=True)
        truth = {"type": "resume", "template": template, "fields": truth_for(person, template),
                 "forbidden_present": ({k: v for k, v in fb.items() if not k.startswith("_")} if template == "t2" else {})}
        (d / "truth.json").write_text(json.dumps(truth, ensure_ascii=False, indent=2), encoding="utf-8")
        render_pdf(doc, d / "doc.pdf")
        render_docx(doc, d / "doc.docx")
        for pi, page in enumerate(render_images(doc), start=1):
            scan_effect(page, rng).save(d / f"scan_{pi}.png", optimize=True)
        manifest.append({"id": doc_id, "template": template, "careers": len(person["career"]), "pages": len(list(d.glob("scan_*.png")))})
        print(doc_id, f"경력 {len(person['career'])}", f"{manifest[-1]['pages']}쪽")

    (out_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"➔ {out_dir.relative_to(REPO)} ({len(manifest)}건)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
